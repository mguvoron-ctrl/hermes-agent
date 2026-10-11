"""
Разбор входящего Excel-инвойса (формат поставщика: листы FAC-IN /
Packing list / Инвойс перевод) → структурированные поля + таблица товаров.

Парсим по ПОДПИСЯМ полей (а не по жёстким адресам ячеек), чтобы небольшие
сдвиги в шаблоне не ломали импорт. Если какое-то поле не нашлось — оставляем
пустым, пользователь дозаполнит руками.

Используется в app.py при загрузке файла на инвойс.
"""

from __future__ import annotations
import re
import datetime as _dt

try:
    from openpyxl import load_workbook
except Exception:  # openpyxl не установлен
    load_workbook = None


def _norm(v) -> str:
    """Нормализация текста ячейки: только буквы/цифры, нижний регистр."""
    if v is None:
        return ""
    s = str(v).lower()
    return re.sub(r"[^0-9a-zа-яё]", "", s)


def _fmt(v) -> str:
    """Человекочитаемое значение ячейки."""
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return f"{round(v, 2):.2f}".rstrip("0").rstrip(".")
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def _matrix(ws) -> list[list]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _trim(mat: list) -> list:
    """Отрезает полностью пустые хвостовые строки (у файлов бывает раздутая
    «используемая область» на тысячи пустых строк — из-за неё всё тормозит)."""
    last = -1
    for i, row in enumerate(mat):
        if any(c not in (None, "") for c in row):
            last = i
    return mat[:last + 1]


def _load_mats(path: str) -> list:
    """Загружает книгу ОДИН раз в быстром read_only-режиме и возвращает матрицы листов."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        return [_trim(_matrix(wb[n])) for n in wb.sheetnames]
    finally:
        wb.close()


def _cell(m: list[list], i: int, j: int):
    if 0 <= i < len(m) and 0 <= j < len(m[i]):
        return m[i][j]
    return None


def _find_right(m, keywords, need_digit=False) -> str:
    """Находит ячейку-подпись и возвращает первое непустое значение справа."""
    for i, row in enumerate(m):
        for j, cell in enumerate(row):
            nc = _norm(cell)
            if nc and any(k in nc for k in keywords):
                for jj in range(j + 1, len(row)):
                    val = _fmt(row[jj])
                    if val:
                        if need_digit and not re.search(r"\d", val):
                            break  # не то поле, пробуем следующую подпись
                        return val
    return ""


def _find_below(m, keywords) -> str:
    """Находит подпись и возвращает первое непустое значение НИЖЕ в том же столбце."""
    for i, row in enumerate(m):
        for j, cell in enumerate(row):
            nc = _norm(cell)
            if nc and any(k in nc for k in keywords):
                for ii in range(i + 1, len(m)):
                    val = _fmt(_cell(m, ii, j))
                    if val:
                        return val
    return ""


def _find_cell_text(m, keywords) -> str:
    """Возвращает полный текст первой ячейки, содержащей ключевое слово."""
    for row in m:
        for cell in row:
            if cell and any(k in _norm(cell) for k in keywords):
                return str(cell).strip()
    return ""


# Штрихкод (EAN-13) и страна происхождения внутри описания товара
_BARCODE = re.compile(r"\b(\d{12,14})\b")
_ORIGIN = re.compile(r"(?:origin|страна\s*происхождения)\s*:?\s*(.+)$",
                     re.IGNORECASE | re.DOTALL)


def _split_description(desc: str) -> tuple[str, str, str]:
    """Из описания достаём: наименование, штрихкод, страну происхождения."""
    if not desc:
        return "", "", ""
    text = str(desc)
    barcode = ""
    mb = _BARCODE.search(text)
    if mb:
        barcode = mb.group(1)
    origin = ""
    mo = _ORIGIN.search(text)
    if mo:
        origin = re.sub(r"\s+", " ", mo.group(1)).strip()
    # Наименование = текст до штрихкода / до "origin"
    name = text
    if mb:
        name = text[:mb.start()]
    name = re.split(r"origin|страна\s*происхождения", name, flags=re.IGNORECASE)[0]
    name = re.sub(r"\s+", " ", name).strip()
    return name, barcode, origin


def _parse_items(m_inv, m_rus, m_pack) -> list[dict]:
    """Таблица товаров из FAC-IN + вес/коробки из Packing list + рус. название."""
    # найти строку заголовка таблицы в FAC-IN
    header_i = None
    for i, row in enumerate(m_inv):
        joined = " ".join(_norm(c) for c in row)
        # заголовок: столбец кода (Custom/Customs Code, Код ТН ВЭД, HS Code) + столбец кол-ва
        has_code = ("custom" in joined or "тнвэд" in joined or "hscode" in joined)
        has_qty = ("qty" in joined or "колво" in joined or "quantity" in joined)
        if has_code and has_qty:
            header_i = i
            break
    if header_i is None:
        return []

    def col_of(m, i, keys):
        for j, c in enumerate(m[i]):
            if any(k in _norm(c) for k in keys):
                return j
        return None

    c_hs = col_of(m_inv, header_i, ["customcode", "кодтнвэд", "customscode"])
    c_item = col_of(m_inv, header_i, ["itemcode", "артикул"])
    c_desc = col_of(m_inv, header_i, ["description", "наименование"])
    c_qty = col_of(m_inv, header_i, ["qty", "колво", "quantity"])
    c_price = col_of(m_inv, header_i, ["unitprice", "ценазашт", "цена"])
    c_total = col_of(m_inv, header_i, ["totalprice", "стоимость", "total"])

    # карта веса/коробок из Packing list по артикулу (по заголовку столбцов)
    pack = {}
    if m_pack:
        ph = None
        for i, row in enumerate(m_pack):
            j = " ".join(_norm(c) for c in row)
            has_code = ("itemcode" in j or "артикул" in j)
            has_wb = any(k in j for k in ("netweight", "nettweight", "нетто", "grossweight", "брутто", "boxes", "короб"))
            if has_code and has_wb:
                ph = i
                break
        if ph is not None:
            pc_item = col_of(m_pack, ph, ["itemcode", "артикул"])
            pc_net = col_of(m_pack, ph, ["nettweight", "netweight", "неттовес", "нетто"])
            pc_gross = col_of(m_pack, ph, ["grossweight", "бруттовес", "брутто"])
            pc_box = col_of(m_pack, ph, ["boxes", "короб"])
            for i in range(ph + 1, len(m_pack)):
                code = _fmt(_cell(m_pack, i, pc_item)) if pc_item is not None else ""
                if not code or _norm(code).startswith(("total", "итого", "grand")):
                    continue
                pack[code] = {
                    "net_weight": _fmt(_cell(m_pack, i, pc_net)) if pc_net is not None else "",
                    "gross_weight": _fmt(_cell(m_pack, i, pc_gross)) if pc_gross is not None else "",
                    "box_count": _fmt(_cell(m_pack, i, pc_box)) if pc_box is not None else "",
                }

    # рус. названия по артикулу (по заголовку столбцов листа «Инвойс перевод»)
    rus = {}
    if m_rus:
        rh = None
        for i, row in enumerate(m_rus):
            j = " ".join(_norm(c) for c in row)
            if ("артикул" in j or "itemcode" in j) and ("наименование" in j or "description" in j):
                rh = i
                break
        if rh is not None:
            rc_item = col_of(m_rus, rh, ["артикул", "itemcode"])
            rc_name = col_of(m_rus, rh, ["наименование", "description"])
            for i in range(rh + 1, len(m_rus)):
                code = _fmt(_cell(m_rus, i, rc_item)) if rc_item is not None else ""
                nm = _cell(m_rus, i, rc_name) if rc_name is not None else ""
                if not code or _norm(code).startswith(("total", "итого", "grand")):
                    continue
                if nm:
                    rus[code] = _split_description(nm)[0]

    items = []
    for i in range(header_i + 1, len(m_inv)):
        hs = _fmt(_cell(m_inv, i, c_hs)) if c_hs is not None else ""
        item_code = _fmt(_cell(m_inv, i, c_item)) if c_item is not None else ""
        desc = _cell(m_inv, i, c_desc) if c_desc is not None else ""
        qty = _fmt(_cell(m_inv, i, c_qty)) if c_qty is not None else ""
        # явный конец таблицы — строка итогов
        if _norm(item_code).startswith(("total", "итого", "grand")) or \
           _norm(hs).startswith(("total", "итого", "grand")):
            break
        # строка товара = есть код ТН ВЭД, артикул и числовое количество
        if not hs or not item_code or \
           not re.fullmatch(r"\d+([.,]\d+)?", qty.replace(" ", "")):
            continue
        name_en, barcode, origin = _split_description(desc)
        name = rus.get(item_code) or name_en  # рус. название, иначе англ.
        it = {
            "hs_code": hs, "item_code": item_code, "name": name, "name_en": name_en,
            "barcode": barcode, "origin": origin, "qty": qty,
            "unit_price": _fmt(_cell(m_inv, i, c_price)) if c_price is not None else "",
            "total_price": _fmt(_cell(m_inv, i, c_total)) if c_total is not None else "",
            "net_weight": "", "gross_weight": "", "box_count": "",
        }
        if item_code in pack:
            it.update(pack[item_code])
        items.append(it)
    return items


def _is_camion(mats) -> bool:
    """Это готовый реестр Камьён (а не инвойс поставщика)?"""
    for m in mats:
        for row in m[:30]:
            j = " ".join(_norm(c) for c in row)
            if "customcode" in j and any(k in j for k in
                                         ("cip", "колвосост", "producteur", "designationrus")):
                return True
    return False


def parse_camion(mats) -> dict:
    """Разбор готового/переделанного файла Камьён (реестр из 15 колонок)."""
    # найти лист и строку заголовка
    m = None
    hdr = None
    for cand in mats:
        for i, row in enumerate(cand[:30]):
            j = " ".join(_norm(c) for c in row)
            if "customcode" in j and ("cip" in j or "designationrus" in j or "qty" in j):
                m, hdr = cand, i
                break
        if m is not None:
            break
    if m is None:
        return {"fields": {}, "items": [], "warnings": ["Не найден заголовок Камьёна"]}

    header = m[hdr]

    def col(*keys, exact=None):
        for j, c in enumerate(header):
            n = _norm(c)
            if exact is not None:
                if n == exact:
                    return j
            elif n and any(k in n for k in keys):
                return j
        return None

    c_hs = col("customcode")
    c_item = col("itemcode")
    c_desig = col(exact="designation")
    c_rus = col("designationrus", "рус")
    c_qty = col("qty", "колво")
    c_cip = col("cip", "ценазашт")
    c_total = col("nettotal", "стоимость")
    c_net = col("netweight", "неттовес")
    c_brut = col("brut", "бруттовес")
    c_box = col("boxes", "короб")
    c_inv = col("invoice", "инвойс")
    c_orig = col("origine", "origin", "страна")
    c_ean = col("ean", "штрих")
    c_set = col("колвосост", "составля", "набор")

    # метаданные
    fields = {}
    fields["supplier"] = _find_right(m, ["отправитель", "exporter", "supplier"])
    cno = _find_right(m, ["контракт", "contract"])
    cdt = _find_right(m, ["датаконтракта", "датаконтр"])
    if cno:
        fields["contract"] = "№ " + cno + (" from " + cdt if cdt else "")

    items = []
    inv_no = ""
    for i in range(hdr + 1, len(m)):
        hs = _fmt(_cell(m, i, c_hs)) if c_hs is not None else ""
        item = _fmt(_cell(m, i, c_item)) if c_item is not None else ""
        if not item and not hs:
            continue
        if _norm(item).startswith(("total", "итого", "grand")):
            break
        if not inv_no and c_inv is not None:
            inv_no = _fmt(_cell(m, i, c_inv))
        items.append({
            "hs_code": hs,
            "item_code": item,
            "name": _fmt(_cell(m, i, c_rus)) if c_rus is not None else "",
            "name_en": _fmt(_cell(m, i, c_desig)) if c_desig is not None else "",
            "barcode": re.sub(r"\D", "", _fmt(_cell(m, i, c_ean))) if c_ean is not None else "",
            "origin": _fmt(_cell(m, i, c_orig)) if c_orig is not None else "",
            "qty": _fmt(_cell(m, i, c_qty)) if c_qty is not None else "",
            "unit_price": _fmt(_cell(m, i, c_cip)) if c_cip is not None else "",
            "total_price": _fmt(_cell(m, i, c_total)) if c_total is not None else "",
            "net_weight": _fmt(_cell(m, i, c_net)) if c_net is not None else "",
            "gross_weight": _fmt(_cell(m, i, c_brut)) if c_brut is not None else "",
            "box_count": _fmt(_cell(m, i, c_box)) if c_box is not None else "",
            "set_count": _fmt(_cell(m, i, c_set)) if c_set is not None else "1",
        })
    if inv_no:
        fields["number"] = inv_no
    fields = {k: v for k, v in fields.items() if v}
    return {"fields": fields, "items": items, "warnings": [], "format": "camion"}


def parse_any(path: str) -> dict:
    """Определяет формат файла (инвойс поставщика или готовый Камьён) и разбирает."""
    if load_workbook is None:
        return {"fields": {}, "items": [], "warnings": ["openpyxl не установлен"], "format": "invoice"}
    mats = _load_mats(path)
    if _is_camion(mats):
        return parse_camion(mats)
    r = _parse_invoice_mats(mats)
    r["format"] = "invoice"
    return r


def parse_invoice(path: str) -> dict:
    """Главная функция. Возвращает {'fields': {...}, 'items': [...], 'warnings': [...]}."""
    if load_workbook is None:
        return {"fields": {}, "items": [], "warnings": ["openpyxl не установлен"]}
    return _parse_invoice_mats(_load_mats(path))


def _parse_invoice_mats(mats: list) -> dict:
    def _header_text(m) -> str:
        """Строка-заголовок товарной таблицы (код + кол-во), нормализованная."""
        for row in m[:45]:
            j = " ".join(_norm(c) for c in row)
            code = any(k in j for k in ("custom", "тнвэд", "hscode", "itemcode", "артикул"))
            qty = any(k in j for k in ("qty", "колво", "quantity"))
            if code and qty:
                return j
        return ""

    # Определяем назначение листов ПО СОДЕРЖИМОМУ (имена бывают «Table 2/3/4»)
    m_inv = m_pack = m_rus = None
    for m in mats:
        j = _header_text(m)
        if not j:
            continue
        has_weight = any(k in j for k in ("netweight", "nettweight", "нетто",
                                          "grossweight", "брутто", "boxes", "короб"))
        is_eng = any(k in j for k in ("unitprice", "totalprice", "customcode",
                                      "customscode", "itemcode"))
        is_rus = any(k in j for k in ("наименование", "ценазашт", "стоимость",
                                      "кодтнвэд", "артикул"))
        if has_weight and m_pack is None:
            m_pack = m
        elif is_eng and not is_rus and m_inv is None:
            m_inv = m
        elif is_rus and m_rus is None:
            m_rus = m
    # фолбэки на случай нестандартной структуры
    if m_inv is None:
        m_inv = next((m for m in mats if _header_text(m) and m is not m_pack), mats[0])

    fields: dict = {}
    fields["number"] = _find_right(m_inv, ["nnvoice", "ninvoice", "invoiceno", "номеринвойса"])
    fields["supplier"] = (_find_below(m_inv, ["exporter", "отправитель", "shipper"]) or "")
    # первая строка блока экспортёра = название компании
    if fields["supplier"]:
        fields["supplier"] = fields["supplier"].splitlines()[0].strip()
    fields["inv_date"] = _find_right(m_inv, ["invoicedate", "датаинвойса", "датаinvoice"])
    fields["incoterms"] = _find_right(m_inv, ["incoterms", "условияпоставки"])
    fields["currency"] = _find_right(m_inv, ["invoicecurrency", "ivoicecurrency", "валютаинвойса", "currency", "валюта"])
    fields["amount"] = _find_right(m_inv, ["totaltopay", "ксумме", "коплате", "итогокоплате"])
    fields["total_qty"] = _find_right(m_inv, ["totalofgoods", "итоготоваров"])
    fields["contract"] = _find_cell_text(m_inv, ["contract", "контракт"])
    fields["consignee"] = _find_below(m_inv, ["invoicingaddress", "юридическийадрес", "invoiceaddress"])
    if fields["consignee"]:
        fields["consignee"] = fields["consignee"].splitlines()[0].strip()
    inn_cell = _find_cell_text(m_inv, ["inn", "инн"])
    m_inn = re.search(r"\d{9,12}", inn_cell)
    fields["inn"] = m_inn.group(0) if m_inn else ""

    # из Packing list — вес, паллеты, коробки, номер PL
    if m_pack:
        fields["packing_list_no"] = _find_right(m_pack, ["packinglist", "упаковочныйлист"], need_digit=True)
        fields["net_weight"] = _find_right(m_pack, ["totalnetweight", "неттовес", "вескгнетто"])
        fields["gross_weight"] = _find_right(m_pack, ["totalgrossweight", "бруттовес"])
        fields["boxes"] = _find_right(m_pack, ["totalboxes", "всегокоробок"])
        fields["pallets"] = _find_right(m_pack, ["totalpallets", "всегопаллет"])

    # запасной вариант паллет/коробок из текстовой строки инвойса
    if not fields.get("pallets") or not fields.get("boxes"):
        txt = _find_cell_text(m_inv, ["pallet", "паллет"])
        mp = re.search(r"(\d+)\s*pallet", txt, re.IGNORECASE)
        mb = re.search(r"(\d+)\s*box", txt, re.IGNORECASE)
        if mp and not fields.get("pallets"):
            fields["pallets"] = mp.group(1)
        if mb and not fields.get("boxes"):
            fields["boxes"] = mb.group(1)

    items = _parse_items(m_inv, m_rus, m_pack)

    warnings = []
    if not fields.get("number"):
        warnings.append("Не найден номер инвойса — заполни вручную.")
    if not items:
        warnings.append("Не распознана таблица товаров — проверь формат файла.")

    return {"fields": {k: v for k, v in fields.items() if v},
            "items": items, "warnings": warnings}
