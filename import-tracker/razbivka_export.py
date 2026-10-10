"""
Экспорт «Разбивки наборов» в Excel (формат листа «Разбивка» из Camion.xlsm):
13 колонок, данные набора объединены по N строкам, составляющие — по строке.
"""

from __future__ import annotations
import io

HEADERS = ["Код ТНВЭД/ Custom Code", "Артикул/ Item Code", "Наименование/ Designation",
           "Описание/ Description", "Код составляющих", "Наименование составляющих в наборе",
           "Декларация соответствия", "От", "До", "ЭЗ", "От", "Спирт", "EAN НАБОРА"]
WIDTHS = [14, 13, 32, 32, 14, 40, 26, 11, 11, 12, 11, 7, 15]
COMP = ["comp_code", "comp_name", "decl", "dfrom", "dto", "ez", "efrom", "spirt"]  # колонки 5..12


def build_razbivka_xlsx(kits) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Разбивка"
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for j, h in enumerate(HEADERS, 1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True)
        c.alignment = center

    r = 2
    for k in kits:
        n = len(k["rows"]) or 1
        g1, g2 = r, r + n - 1
        # данные набора в первой строке блока
        ws.cell(r, 1, k["hs_code"])
        ws.cell(r, 2, k["item_code"])
        ws.cell(r, 3, k["name_en"])
        ws.cell(r, 4, k["name"])
        ws.cell(r, 13, k["barcode"]).number_format = "@"
        # составляющие по строкам
        for i, row in enumerate(k["rows"]):
            rr = r + i
            for off, field in enumerate(COMP):
                ws.cell(rr, 5 + off, row.get(field, ""))
        # объединяем столбцы набора по блоку
        if g2 > g1:
            for col in (1, 2, 3, 4, 13):
                ws.merge_cells(start_row=g1, start_column=col, end_row=g2, end_column=col)
                ws.cell(g1, col).alignment = center
        r += n

    last = r - 1
    for j, w in enumerate(WIDTHS, 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(j)].width = w
    if last >= 1:
        for row in ws.iter_rows(min_row=1, max_row=last, min_col=1, max_col=13):
            for c in row:
                c.border = border
                if not c.alignment.wrap_text:
                    c.alignment = Alignment(vertical="center", wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
