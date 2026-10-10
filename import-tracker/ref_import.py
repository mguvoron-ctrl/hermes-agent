"""
Импорт справочников TRANSIT и НОВИНКИ → таблица article_ref.

По каждому артикулу храним актуальные данные:
  - из TRANSIT берём строку САМОЙ СВЕЖЕЙ поставки (Supply вида "26/076");
  - артикулы, которых нет в TRANSIT, добираем из файлов НОВИНКИ.

Используется со страницы «Справочники».
"""

from __future__ import annotations
import re
import tempfile
import os

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


def norm_art(s) -> str:
    """Артикул без пробелов/дефисов/тире, в верхнем регистре (как NormArt в VBA)."""
    if s is None:
        return ""
    out = []
    for ch in str(s):
        if ord(ch) in (32, 45, 9, 10, 13, 160, 8211, 8212):
            continue
        out.append(ch)
    return "".join(out).upper()


def _norm_hdr(v) -> str:
    return re.sub(r"\s+", "", str(v or "").lower())


def supply_key(s) -> int:
    """"26/076" -> 26076 (год*1000 + номер); иначе -1."""
    m = re.match(r"\s*(\d+)\s*/\s*(\d+)", str(s or ""))
    if not m:
        return -1
    return int(m.group(1)) * 1000 + int(m.group(2))


def _clean(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").replace("\xa0", " ")).strip()


# --- Матрица листа для xlsx и xlsb -----------------------------------------
def _sheets_xlsx(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        yield sn, ws.iter_rows(values_only=True)


def _sheets_xlsb(path):
    from pyxlsb import open_workbook
    with open_workbook(path) as wb:
        for sn in wb.sheets:
            with wb.get_sheet(sn) as sh:
                yield sn, ([c.v for c in row] for row in sh.rows())


def _iter_sheets(path):
    if path.lower().endswith(".xlsb"):
        yield from _sheets_xlsb(path)
    else:
        yield from _sheets_xlsx(path)


def _find_col(header_row, keywords) -> int | None:
    for j, cell in enumerate(header_row):
        h = _norm_hdr(cell)
        if h and any(k in h for k in keywords):
            return j
    return None


# --- TRANSIT ---------------------------------------------------------------
def parse_transit(path: str) -> list[dict]:
    """Возвращает по строке на артикул — из самой свежей поставки."""
    if load_workbook is None:
        return []
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    c_sup = _find_col(header, ["supply", "поставк"])
    c_art = _find_col(header, ["артикул", "itemcode"])
    c_hs = _find_col(header, ["тнвэд", "customcode"])
    c_ean = _find_col(header, ["штрих", "ean"])
    c_rus = _find_col(header, ["рус"])
    c_orig = _find_col(header, ["origine", "origin"])
    c_prod = _find_col(header, ["fournisseur", "producteur", "производ"])

    def g(row, c):
        return row[c] if (c is not None and c < len(row)) else None

    best_sk: dict[str, int] = {}
    best: dict[str, dict] = {}
    for row in rows:
        art = norm_art(g(row, c_art))
        if not art:
            continue
        sk = supply_key(g(row, c_sup))
        if art not in best_sk or sk > best_sk[art]:
            best_sk[art] = sk
            best[art] = {
                "article": art,
                "hs_code": _clean(g(row, c_hs)),
                "ean": re.sub(r"\D", "", str(g(row, c_ean) or "")),
                "name_rus": _clean(g(row, c_rus)),
                "origin": _clean(g(row, c_orig)),
                "producer": _clean(g(row, c_prod)),
                "supply_key": sk,
            }
    return list(best.values())


# --- НОВИНКИ ---------------------------------------------------------------
def parse_novinki(path: str) -> list[dict]:
    """Артикул -> код/EAN/рус. название из файла новинок (по всем листам)."""
    out: dict[str, dict] = {}
    for sn, rowiter in _iter_sheets(path):
        rowlist = list(rowiter)
        # заголовок: среди первых 10 строк — есть "артикул" и (тнвэд или ean)
        hdr_i = None
        for i, row in enumerate(rowlist[:10]):
            j = " ".join(_norm_hdr(c) for c in row)
            if "артикул" in j and ("тнвэд" in j or "ean" in j or "русское" in j):
                hdr_i = i
                break
        if hdr_i is None:
            continue
        header = rowlist[hdr_i]
        c_art = _find_col(header, ["артикул"])
        c_hs = _find_col(header, ["тнвэд"])
        c_ean = _find_col(header, ["ean"])
        c_rus = _find_col(header, ["русское", "русскоеназвание"])
        if c_art is None:
            continue

        def g(row, c):
            return row[c] if (c is not None and c < len(row)) else None

        for row in rowlist[hdr_i + 1:]:
            art = norm_art(g(row, c_art))
            if not art or art in out:
                continue
            hs = _clean(g(row, c_hs))
            ean = re.sub(r"\D", "", str(g(row, c_ean) or ""))
            rus = _clean(g(row, c_rus))
            if not (hs or ean or rus):
                continue
            out[art] = {"article": art, "hs_code": hs, "ean": ean,
                        "name_rus": rus, "origin": "", "producer": "", "supply_key": 0}
    return list(out.values())
