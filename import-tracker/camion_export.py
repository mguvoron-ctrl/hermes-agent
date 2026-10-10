"""
Сборка готового Камьён-реестра в Excel (метаданные + 15 колонок + подсветка
проблем сверки). Вынесено из app.py, чтобы маршруты остались тонкими.

build_camion_xlsx(invoice_row, vitems) -> bytes
  invoice_row — строка инвойса из БД;
  vitems — список товаров со статусами сверки (из validate_items).
"""

from __future__ import annotations
import io
import re


def xl_num(v):
    """Строку-число -> число (для Excel); иначе исходная строка; пусто -> ''."""
    s = str(v or "").strip().replace(" ", "")
    if s == "":
        return ""
    try:
        f = float(s.replace(",", "."))
        return int(f) if f == int(f) else f
    except ValueError:
        return str(v).strip()


HEADERS = ["Custom code", "Item code", "Designation", "Designation rus", "Qty",
           "CIP", "Net total", "Net Weight, kg", "Brut Weigt, kg", "Boxes",
           "Invoice", "Origine", "Producteur", "EAN", "Кол-во сост."]
WIDTHS = [13, 14, 34, 42, 8, 8, 11, 13, 13, 7, 15, 13, 24, 16, 9]


def build_camion_xlsx(invoice, vitems) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"
    fill = {"bad": PatternFill("solid", fgColor="FF0000"),
            "warn": PatternFill("solid", fgColor="FFFF00")}

    contract = invoice["contract"] or ""
    mno = re.search(r"№\s*([\w./\-]+)", contract)
    mdt = re.search(r"(?:from|от)\s*([\d./\-]+)", contract, re.I)
    ws["A1"] = "Отправитель:"; ws["B1"] = invoice["supplier"] or ""
    ws["A2"] = "Контракт №:"; ws["B2"] = mno.group(1) if mno else ""
    ws["A3"] = "Дата контракта:"; ws["B3"] = mdt.group(1) if mdt else ""

    for j, h in enumerate(HEADERS, 1):
        ws.cell(4, j, h).font = Font(bold=True)

    r = 5
    for it in vitems:
        ws.cell(r, 1, xl_num(it["hs_code"]))
        ws.cell(r, 2, it["item_code"])
        ws.cell(r, 3, it["name_en"])
        ws.cell(r, 4, it["name"])
        ws.cell(r, 5, xl_num(it["qty"]))
        ws.cell(r, 6, xl_num(it["unit_price"]))
        ws.cell(r, 7, xl_num(it["total_price"]))
        ws.cell(r, 8, xl_num(it["net_weight"]))
        ws.cell(r, 9, xl_num(it["gross_weight"]))
        ws.cell(r, 10, xl_num(it["box_count"]))
        ws.cell(r, 11, invoice["number"])
        ws.cell(r, 12, it["origin"])
        ws.cell(r, 13, it["producer"])
        ws.cell(r, 14, it["barcode"]).number_format = "@"
        ws.cell(r, 15, 1)
        for col, st in ((1, it["hs_st"]), (4, it["rus_st"]), (12, it["origin_st"]),
                        (13, it["prod_st"]), (14, it["ean_st"])):
            if st in fill:
                ws.cell(r, col).fill = fill[st]
        r += 1

    for j, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
